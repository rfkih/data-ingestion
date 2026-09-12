"""Parser for IDX's standardized financial-statement workbook (FinancialStatement-<year>-<period>-<code>.xlsx).

Layout (every taxonomy family — general 2xxxxxx, infrastructure 3xxxxxx, banks 4xxxxxx, ...):
  sheet ``1000000``  general information (sector, board, period dates, rounding unit, currency)
  numeric sheets     one statement each; row 1 = ``[code] title``, row 3 = context names per value column
                     (``CurrentYearInstant``/``PriorEndYearInstant`` or ``CurrentYearDuration``/``PriorYearDuration``),
                     data rows = ``[indonesian label, value..., english label]`` read POSITIONALLY (a missing value
                     is an empty cell, never a shifted column).
Values are scaled to full IDR using the "level of rounding" row. English labels are the concept keys.
"""
from __future__ import annotations

import logging
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger(__name__)

ROUNDING = {"satuan penuh": 1, "full amount": 1, "ribuan": 1_000, "thousand": 1_000, "jutaan": 1_000_000,
            "million": 1_000_000, "miliaran": 1_000_000_000, "billion": 1_000_000_000}
STATEMENT_TYPES = {"financial position": "position", "profit or loss": "pnl", "cash flows": "cashflow",
                   "changes in equity": "equity", "general information": "info"}


@dataclass
class Fact:
    sheet: str
    row: int
    concept: str          # English label
    context: str          # CurrentYearInstant, PriorYearDuration, ...
    value: Decimal        # scaled to full IDR (per-share values are NOT scaled)


@dataclass
class Parsed:
    info: dict[str, Any] = field(default_factory=dict)
    rounding: int = 1                   # as labelled on the info sheet
    rounding_eff: int = 1               # as applied (label overridden when the values are plainly already full amounts)
    currency: str | None = None
    fx_rate: Decimal | None = None      # IDR per unit of presentation currency (reporting-date rate from the info sheet)
    fx_source: str | None = None        # info_sheet | fallback | None
    statements: dict[str, str] = field(default_factory=dict)     # sheet -> statement type
    facts: list[Fact] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# Plausibility bounds for "total assets" after scaling. No IDR filer has 50,000 T of assets (the largest bank has ~1,700 T)
# and no USD filer has 3 trillion USD; a labelled-in-millions workbook whose values are already full amounts breaks these.
# The FY2023 workbook vintage (published Jan-Apr 2024) did exactly that for most large filers. Below 1 bn rupiah is equally
# impossible for a listed company: a "full amount" label over values that are really in millions.
MAX_ASSETS = {"IDR": Decimal(50_000_000_000_000_000), "USD": Decimal(3_000_000_000_000)}
MIN_ASSETS_IDR = Decimal(1_000_000_000)


def _num(v: Any) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip().replace(",", "")
    if not s or s in ("-", "—"):
        return None
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def _date(v: Any) -> date | None:
    if v is None:
        return None
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


PER_SHARE = re.compile(r"per share", re.I)
_DATE_HDR = re.compile(r"^\d{1,2} \w+ \d{4}$|^\d{4}-\d{2}-\d{2}")


def _is_context_header(c: Any) -> bool:
    """2023+ workbooks label value columns CurrentYearInstant/PriorYearDuration...; FY2020-2022 workbooks
    put the period-end date there ('31 December 2020'). Both mark a context column."""
    if isinstance(c, (datetime, date)):
        return True
    t = str(c).strip()
    return bool(re.match(r"^(Current|Prior)", t) or _DATE_HDR.match(t))
CURRENT = {"position": "CurrentYearInstant", "pnl": "CurrentYearDuration", "cashflow": "CurrentYearDuration",
           "equity": "CurrentYearDuration"}
PRIOR = {"position": "PriorEndYearInstant", "pnl": "PriorYearDuration", "cashflow": "PriorYearDuration",
         "equity": "PriorYearDuration"}


def parse_workbook(path: str | Path, fx_fallback: Callable[[date], Decimal | None] | None = None,
                   rounding_override: int | None = None) -> Parsed:
    """``fx_fallback(period_end)`` supplies the IDR rate for a non-IDR filer whose info sheet lacks (or garbles) it;
    ``rounding_override`` forces the scale (the store uses it after cross-checking total assets with the neighbouring report)."""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    p = Parsed()
    if "1000000" in wb.sheetnames:
        _parse_info(wb["1000000"], p)
    if p.currency and p.currency != "IDR":
        if p.fx_rate is not None and p.fx_rate < 1000:
            p.warnings.append(f"conversion rate {p.fx_rate} on the info sheet is not an IDR rate; ignored")
            p.fx_rate = None
        if p.fx_rate is not None:
            p.fx_source = "info_sheet"
        elif fx_fallback is not None and p.info.get("period_end"):
            p.fx_rate = fx_fallback(p.info["period_end"])
            p.fx_source = "fallback" if p.fx_rate else None
        if p.fx_rate is None:
            p.warnings.append(f"{p.currency} filer without a conversion rate: values left in {p.currency}")
    raw_facts: list[tuple[str, int, str, str, Decimal, bool]] = []
    for name in wb.sheetnames:
        if not name[:1].isdigit() or name == "1000000" or name.endswith("PY"):
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        if len(rows) < 4:
            continue
        title = " ".join(str(v) for v in rows[0] if v) if rows[0] else ""
        stype = next((t for k, t in STATEMENT_TYPES.items() if k in title.lower()), None)
        if stype is None:
            continue
        p.statements[name] = stype
        ctx_idx, contexts = None, {}
        for j in range(1, min(8, len(rows))):
            cand = {i: str(c).strip() for i, c in enumerate(rows[j])
                    if i > 0 and c is not None and _is_context_header(c)}
            if cand and rows[j][0] in (None, ""):          # header rows carry no label in column A
                ctx_idx, contexts = j, cand
                break
        if not contexts:
            continue
        # IDX's header labels are inconsistent across filers (e.g. "CurrentYearInstant" on a cash-flow sheet);
        # the semantics are positional: first value column = current period, second = prior. Normalise.
        want_cur, want_pri = CURRENT.get(stype, "CurrentYear"), PRIOR.get(stype, "PriorYear")
        ordered = sorted(contexts)
        norm = {}
        for k, col in enumerate(ordered):
            raw = contexts[col]
            if raw.startswith("Current") or (k == 0 and not raw.startswith("Prior")):
                norm[col] = want_cur
            elif raw.startswith("Prior") or k == 1:
                norm[col] = want_pri
            else:
                norm[col] = raw
        contexts = norm
        last_col = max(contexts)
        for r_idx, row in enumerate(rows[ctx_idx + 1:], start=ctx_idx + 2):
            if not row or len(row) <= last_col:
                continue
            eng = None
            for cell in reversed(row[last_col + 1:]):
                if cell is not None and str(cell).strip():
                    eng = str(cell).strip()
                    break
            if not eng:
                continue
            per_share = bool(PER_SHARE.search(eng))
            for col, ctx in contexts.items():
                v = _num(row[col])
                if v is None:
                    continue
                raw_facts.append((name, r_idx, eng, ctx, v, per_share))
    wb.close()
    p.rounding_eff = rounding_override if rounding_override is not None else _effective_rounding(p, raw_facts)
    if rounding_override is not None and rounding_override != p.rounding:
        p.warnings.append(f"rounding label '{p.info.get('rounding_label')}' overridden to x{rounding_override} from the neighbouring report")
    fx = p.fx_rate if (p.currency and p.currency != "IDR" and p.fx_rate) else Decimal(1)
    for name, r_idx, eng, ctx, v, per_share in raw_facts:
        p.facts.append(Fact(name, r_idx, eng, ctx, v * (fx if per_share else p.rounding_eff * fx)))
    return p


def _effective_rounding(p: Parsed, raw: list[tuple[str, int, str, str, Decimal, bool]]) -> int:
    """The label says 'in millions' but the numbers are already full amounts (or the reverse): decide from total assets."""
    ta = next((v for _, _, eng, ctx, v, _ in raw if eng.lower() == "total assets" and ctx.startswith("Current") and v > 0), None)
    if ta is None:
        return p.rounding
    cur = "USD" if (p.currency and p.currency != "IDR") else "IDR"
    scaled = ta * p.rounding
    if scaled > MAX_ASSETS[cur] and p.rounding > 1:
        p.warnings.append(f"rounding label '{p.info.get('rounding_label')}' ignored: total assets {ta:,.0f} are already full amounts")
        return 1
    if cur == "IDR" and p.rounding == 1 and scaled < MIN_ASSETS_IDR:
        p.warnings.append(f"rounding label '{p.info.get('rounding_label')}' ignored: total assets {ta:,.0f} must be in millions")
        return 1_000_000
    return p.rounding


def _parse_info(ws, p: Parsed) -> None:
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v is not None and str(v).strip()]
        if len(vals) < 2:
            continue
        key = str(vals[-1]).strip().lower()
        val = vals[1] if len(vals) >= 3 else None
        if val is None:
            continue
        sval = str(val).strip()
        if key.startswith("entity code"):
            p.info["code"] = sval.upper()
        elif key == "sector":
            p.info["sector"] = sval
        elif key == "subsector":
            p.info["subsector"] = sval
        elif key == "industry":
            p.info["industry"] = sval
        elif key.startswith("type of board"):
            p.info["board"] = sval
        elif key.startswith("current period start"):
            p.info["period_start"] = _date(val)
        elif key.startswith("current period end"):
            p.info["period_end"] = _date(val)
        elif key.startswith("prior period end") or key.startswith("prior year end"):
            p.info["prior_period_end"] = _date(val)
        elif key.startswith("period of financial statements"):
            p.info["period_label"] = sval
        elif key.startswith("description of presentation currency"):
            low = sval.lower()
            p.currency = ("IDR" if ("idr" in low or "rupiah" in low) else "USD" if ("usd" in low or "dollar" in low or "dolar" in low)
                          else sval[:12])
        elif key.startswith("conversion rate at reporting date"):
            rate = _num(sval.replace(".", "").replace(",", ".")) if sval.count(".") == 1 and len(sval.split(".")[1]) == 3 else _num(sval)
            if rate is not None and rate > 0:
                p.fx_rate = rate
        elif key.startswith("level of rounding"):
            p.rounding = next((m for k, m in ROUNDING.items() if k in sval.lower()), 1)
            p.info["rounding_label"] = sval
        elif key.startswith("entity main industry"):
            p.info["main_industry"] = sval
        elif key.startswith("whether the financial statements are of an individual entity or a group"):
            p.info["group"] = sval


# ---------------------------------------------------------------------------
# concept map -> metrics (current period; prior period from the same report)
# ---------------------------------------------------------------------------
# (statement, [exact English labels in priority order]) ; SUM = add every present label
POSITION = "position"
PNL = "pnl"
CF = "cashflow"
METRICS: dict[str, tuple[str, list[str], str]] = {
    "total_assets":      (POSITION, ["Total assets"], "first"),
    "total_liabilities": (POSITION, ["Total liabilities"], "first"),
    "total_equity":      (POSITION, ["Total equity"], "first"),
    "equity_parent":     (POSITION, ["Total equity attributable to equity owners of parent entity"], "first"),
    "cash":              (POSITION, ["Cash and cash equivalents", "Cash"], "first"),
    "current_assets":    (POSITION, ["Total current assets"], "first"),
    "current_liabilities": (POSITION, ["Total current liabilities"], "first"),
    "total_debt":        (POSITION, ["Short term bank loans", "Short-term non-bank loans", "Current maturities of bank loans",
                                     "Current maturities of finance lease liabilities", "Current maturities of medium term notes",
                                     "Current maturities of bonds payable", "Current maturities of sukuk",
                                     "Current maturities of other borrowings", "Long-term bank loans",
                                     "Long-term finance lease liabilities", "Long-term medium term notes",
                                     "Long-term bonds payable", "Long-term sukuk", "Long-term other borrowings",
                                     "Borrowings third parties", "Borrowings related parties", "Bonds payable", "Sukuk",
                                     "Medium term notes", "Subordinated loans", "Subordinated bonds"], "sum"),
    "revenue":           (PNL, ["Sales and revenue", "Interest income", "Revenue from insurance premiums",
                                "Revenue from consumer financing"], "first"),
    "gross_profit":      (PNL, ["Total gross profit"], "first"),
    "pretax_profit":     (PNL, ["Total profit (loss) before tax"], "first"),
    "net_profit_total":  (PNL, ["Total profit (loss)"], "first"),
    "net_profit":        (PNL, ["Profit (loss) attributable to parent entity"], "first"),
    "finance_cost":      (PNL, ["Interest and finance costs"], "first"),
    "eps":               (PNL, ["Basic earnings (loss) per share from continuing operations",
                                "Basic earnings per share attributable to equity owners of the parent entity"], "first"),
    "cfo":               (CF, ["Total net cash flows received from (used in) operating activities"], "first"),
    "capex":             (CF, ["Payments for acquisition of property and equipment",
                               "Payments for acquisition of property, plant and equipment",
                               "Acquisition of property, plant and equipment"], "first"),
    "dividends_paid":    (CF, ["Dividends paid from financing activities", "Dividends paid from operating activities"], "sum"),
}


def extract_metrics(p: Parsed) -> dict[str, dict[str, Decimal | None]]:
    """{'current': {metric: value}, 'prior': {...}} using exact English labels; first numeric row wins."""
    by = {}
    for f in p.facts:
        st = p.statements.get(f.sheet)
        if st in ("position", "pnl", "cashflow"):
            by.setdefault((st, f.concept, f.context), []).append(f)
    out = {"current": {}, "prior": {}}
    for metric, (st, labels, mode) in METRICS.items():
        for which, ctxmap in (("current", CURRENT), ("prior", PRIOR)):
            ctx = ctxmap[st]
            total = Decimal(0)
            found = False
            first_val = None
            for lab in labels:
                fs = by.get((st, lab, ctx))
                if not fs:
                    continue
                v = fs[0].value
                if mode == "first":
                    first_val = v
                    break
                total += v
                found = True
            out[which][metric] = first_val if mode == "first" else (total if found else None)
    return out
