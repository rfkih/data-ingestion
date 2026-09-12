"""Workbook parser: both IDX header styles, unit scaling, USD conversion, positional columns."""
from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl

from blackheart_ingest.idx.fin_parse import _is_context_header, extract_metrics, parse_workbook


def _book(tmp_path, header, rounding="Satuan Penuh / Full Amount", currency="Rupiah / IDR", rate=None, assets=2_000_000_000):
    wb = openpyxl.Workbook()
    info = wb.active
    info.title = "1000000"
    info.append(["[1000000] General information"])
    info.append([])
    info.append(["Informasi umum", None, "General information"])
    info.append(["Kode entitas", "TEST", "Entity code"])
    info.append(["Sektor", "B. Basic Materials", "Sector"])
    info.append(["Tanggal awal periode berjalan", "2024-01-01", "Current period start date"])
    info.append(["Tanggal akhir periode berjalan", "2024-12-31", "Current period end date"])
    info.append(["Mata uang pelaporan", currency, "Description of presentation currency"])
    if rate:
        info.append(["Kurs konversi", rate, "Conversion rate at reporting date if presentation currency is other than rupiah"])
    info.append(["Pembulatan", rounding, "Level of rounding used in financial statements"])
    pos = wb.create_sheet("2210000")
    pos.append(["[2210000] Statement of financial position"])
    pos.append([])
    pos.append(["Laporan posisi keuangan", None, "Statement of financial position"])
    pos.append([None, *header])
    pos.append(["Kas dan setara kas", 100, 80, "Cash and cash equivalents"])
    pos.append(["Jumlah aset", assets, None, "Total assets"])              # prior missing -> must not shift
    pos.append(["Jumlah ekuitas induk", 400, 350, "Total equity attributable to equity owners of parent entity"])
    pnl = wb.create_sheet("2311000")
    pnl.append(["[2311000] Statement of profit or loss"])
    pnl.append([])
    pnl.append(["Laba rugi", None, "Statement of profit or loss"])
    pnl.append([None, *header])
    pnl.append(["Penjualan", 500, 450, "Sales and revenue"])
    pnl.append(["Laba induk", 50, 40, "Profit (loss) attributable to parent entity"])
    pnl.append(["LPS", 5.5, 4.4, "Basic earnings (loss) per share from continuing operations"])
    p = tmp_path / "t.xlsx"
    wb.save(p)
    return p


def test_context_header_styles() -> None:
    assert _is_context_header("CurrentYearInstant") and _is_context_header("PriorYearDuration")
    assert _is_context_header("31 December 2020") and _is_context_header("2021-06-30")
    assert not _is_context_header("Statement of financial position") and not _is_context_header("Aset")


def test_new_style_headers_scale_and_positional(tmp_path) -> None:
    p = parse_workbook(_book(tmp_path, ["CurrentYearInstant", "PriorEndYearInstant"], rounding="Jutaan / In Million"))
    m = extract_metrics(p)
    assert p.rounding == 1_000_000 and p.info["period_end"] == date(2024, 12, 31)
    assert m["current"]["total_assets"] == Decimal(2_000_000_000) * 1_000_000 and m["prior"]["total_assets"] is None
    assert m["current"]["cash"] == Decimal(100) * 1_000_000 and m["prior"]["cash"] == Decimal(80) * 1_000_000
    assert m["current"]["eps"] == Decimal("5.5")                       # per-share values are not scaled


def test_old_style_date_headers(tmp_path) -> None:
    p = parse_workbook(_book(tmp_path, ["31 December 2020", "31 December 2019"]))
    m = extract_metrics(p)
    assert len(p.facts) == 11
    assert m["current"]["net_profit"] == Decimal(50) and m["prior"]["net_profit"] == Decimal(40)
    assert m["current"]["revenue"] == Decimal(500)


def test_usd_filer_converted_at_reporting_rate(tmp_path) -> None:
    p = parse_workbook(_book(tmp_path, ["CurrentYearInstant", "PriorEndYearInstant"], rounding="Ribuan / In Thousand",
                             currency="Dollar Amerika / USD", rate="17.856"))
    m = extract_metrics(p)
    assert p.currency == "USD" and p.fx_rate == Decimal(17856)
    assert m["current"]["cash"] == Decimal(100) * 1000 * 17856
    assert m["current"]["eps"] == Decimal("5.5") * 17856                # USD/share -> IDR/share


def test_rounding_label_overridden_when_values_are_already_full(tmp_path) -> None:
    # the FY2023 vintage: labelled "in millions" but total assets 1,408,107,010,000,000 are plainly full rupiah
    p = parse_workbook(_book(tmp_path, ["CurrentYearInstant", "PriorEndYearInstant"], rounding="Jutaan / In Million",
                             assets=1_408_107_010_000_000))
    m = extract_metrics(p)
    assert p.rounding == 1_000_000 and p.rounding_eff == 1
    assert m["current"]["total_assets"] == Decimal(1_408_107_010_000_000)
    assert m["current"]["cash"] == Decimal(100)                        # every value follows the same decision
    assert any("already full amounts" in w for w in p.warnings)


def test_full_amount_label_over_millions_is_scaled_up(tmp_path) -> None:
    p = parse_workbook(_book(tmp_path, ["CurrentYearInstant", "PriorEndYearInstant"], assets=250_000))   # Rp 250 bn in millions
    assert p.rounding_eff == 1_000_000 and extract_metrics(p)["current"]["total_assets"] == Decimal(250_000) * 1_000_000


def test_usd_filer_without_rate_uses_fallback(tmp_path) -> None:
    p = parse_workbook(_book(tmp_path, ["CurrentYearInstant", "PriorEndYearInstant"], currency="Dollar Amerika / USD"),
                       fx_fallback=lambda d: Decimal(16000) if d == date(2024, 12, 31) else None)
    assert p.fx_source == "fallback" and p.fx_rate == Decimal(16000)
    assert extract_metrics(p)["current"]["cash"] == Decimal(100) * 16000
    q = parse_workbook(_book(tmp_path, ["CurrentYearInstant", "PriorEndYearInstant"], currency="Dollar Amerika / USD"))
    assert q.fx_rate is None and any("without a conversion rate" in w for w in q.warnings)


def test_scale_fix_from_neighbouring_report() -> None:
    from blackheart_ingest.idx.fin_store import _scale_fix
    # labelled thousands (x1000) but the numbers were full: 9,911 T vs 8 T last year -> undo the label
    assert _scale_fix(Decimal(9_911_254_000_000_000), Decimal(8_049_000_000_000), 1000) == 1
    # labelled millions, values in millions, neighbour agrees -> nothing to do
    assert _scale_fix(Decimal(1_449_000_000_000_000), Decimal(1_408_000_000_000_000), 1_000_000) is None
    # "full amount" label over values that are really in thousands
    assert _scale_fix(Decimal(204_000_000), Decimal(190_000_000_000), 1) == 1000
    # a real 97x jump (reverse takeover) is not a power of 1000 -> untouched
    assert _scale_fix(Decimal(15_938_000_000_000), Decimal(164_000_000_000), 1000) is None
    # 3,489x = a lying label on top of real 3.5x growth -> still a power of 1000 within tolerance
    assert _scale_fix(Decimal(251_200_000_000_000), Decimal(72_000_000_000), 1000) == 1
    assert _scale_fix(None, Decimal(1), 1000) is None and _scale_fix(Decimal(1), Decimal(0), 1000) is None


def test_eps_reconcile_uses_the_price_to_pick_the_wrong_side() -> None:
    from blackheart_ingest.idx.fin_store import _eps_reconcile
    # BBNI FY2024: profit right (Rp 21.5 T), EPS stored 0.0006 (the filer scaled the per-share row too) -> EPS x 1e6
    o, eps, flags = _eps_reconcile(Decimal(21_463_599_000_000), Decimal("0.0006"), 36_924_339_786, Decimal(4500), 1_000_000)
    assert o is None and eps == Decimal(600) and flags == ["eps_rescaled"]
    # PGEO FY2024: every figure 1000x too big under a "millions" label, EPS (Rp 62) fine -> re-parse the report at x1000
    o, eps, flags = _eps_reconcile(Decimal(2_593_101_558_000_000), Decimal(62), 41_508_024_149, Decimal(900), 1_000_000)
    assert o == 1000 and eps == Decimal(62) and flags == ["report_rescaled"]
    # BRMS FY2023: a high-P/E stock (profit Rp 215 bn, 142 bn shares, price 150); EPS stored 1541.6 = 1000x -> EPS / 1000
    o, eps, flags = _eps_reconcile(Decimal(214_582_441_608), Decimal("1541.6"), 141_784_040_338, Decimal(150), 1)
    assert o is None and eps == Decimal("1.5416") and flags == ["eps_rescaled"]
    # consistent -> nothing; off by something that is not a power of 1000 -> eps_mismatch only; no price -> undecidable
    assert _eps_reconcile(Decimal(1000), Decimal(10), 100, Decimal(50), 1) == (None, Decimal(10), [])
    assert _eps_reconcile(Decimal(1000), Decimal(2), 100, Decimal(50), 1)[2] == ["eps_mismatch"]
    assert _eps_reconcile(Decimal(1_000_000), Decimal(10), 100, None, 1)[2] == ["scale_mismatch"]
    assert _eps_reconcile(None, Decimal(10), 100, Decimal(50), 1) == (None, Decimal(10), [])
